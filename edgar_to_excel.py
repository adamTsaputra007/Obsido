

#!/usr/bin/env python3
"""
Pull data from SEC EDGAR (official APIs) and save it to an Excel workbook.

Install dependencies:
    pip install requests pandas openpyxl
"""

import argparse
import os
import sys
import time
from datetime import date

import pandas as pd
import requests
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TICKER_URL = "https://www.sec.gov/files/company_tickers.json"
SUBMISSIONS_URL = "https://data.sec.gov/submissions/CIK{cik}.json"
FACTS_URL = "https://data.sec.gov/api/xbrl/companyfacts/CIK{cik}.json"
FILING_URL = "https://www.sec.gov/Archives/edgar/data/{cik_int}/{accn_nodash}/{doc}"

METRICS = {
    "Revenue": (["Revenues", "RevenueFromContractWithCustomerExcludingAssessedTax",
                 "SalesRevenueNet"], "USD"),
    "Gross Profit": (["GrossProfit"], "USD"),
    "Operating Income": (["OperatingIncomeLoss"], "USD"),
    "Net Income": (["NetIncomeLoss", "ProfitLoss"], "USD"),
    "Diluted EPS": (["EarningsPerShareDiluted"], "USD/shares"),
    "Cash & Equivalents": (["CashAndCashEquivalentsAtCarryingValue"], "USD"),
    "Total Assets": (["Assets"], "USD"),
    "Total Liabilities": (["Liabilities"], "USD"),
    "Shareholders' Equity": (["StockholdersEquity"], "USD"),
    "Operating Cash Flow": (["NetCashProvidedByUsedInOperatingActivities"], "USD"),
}

FONT = "Arial"


class Edgar:
    def __init__(self, user_agent):
        self.session = requests.Session()
        self.session.headers.update({"User-Agent": user_agent,
                                     "Accept-Encoding": "gzip, deflate"})
        self._tickers = None

    def get_json(self, url):
        time.sleep(0.15)
        for attempt in range(1, 4):
            try:
                r = self.session.get(url, timeout=30)
                if r.status_code == 404:
                    return None
                r.raise_for_status()
                return r.json()
            except requests.RequestException as e:
                print(f"[warn] {url} (attempt {attempt}/3): {e}", file=sys.stderr)
                time.sleep(2 ** attempt)
        return None

    def lookup(self, ticker):
        if self._tickers is None:
            data = self.get_json(TICKER_URL) or {}
            self._tickers = {v["ticker"].upper(): v for v in data.values()}
        hit = self._tickers.get(ticker.upper())
        if not hit:
            return None, None
        return str(hit["cik_str"]).zfill(10), hit["title"]

    def filings(self, cik, forms, count):
        data = self.get_json(SUBMISSIONS_URL.format(cik=cik))
        if not data:
            return []
        recent = data["filings"]["recent"]
        cik_int = int(cik)
        rows = []
        for i, form in enumerate(recent["form"]):
            if forms and form not in forms:
                continue
            accn = recent["accessionNumber"][i]
            rows.append({
                "Form": form,
                "Filing Date": recent["filingDate"][i],
                "Report Date": recent["reportDate"][i] or None,
                "Description": recent.get("primaryDocDescription", [""] * (i + 1))[i],
                "Accession No.": accn,
                "Link": FILING_URL.format(cik_int=cik_int,
                                          accn_nodash=accn.replace("-", ""),
                                          doc=recent["primaryDocument"][i]),
            })
            if len(rows) >= count:
                break
        return rows

    def annual_financials(self, cik, years):
        data = self.get_json(FACTS_URL.format(cik=cik))
        if not data:
            return None
        gaap = data.get("facts", {}).get("us-gaap", {})

        table = {}
        for label, (tags, unit) in METRICS.items():
            for tag in tags:
                items = gaap.get(tag, {}).get("units", {}).get(unit, [])
                by_end = {}
                for it in items:
                    if it.get("form") != "10-K":
                        continue
                    if "start" in it:
                        days = (date.fromisoformat(it["end"]) -
                                date.fromisoformat(it["start"])).days
                        if not 350 <= days <= 380:
                            continue
                    prev = by_end.get(it["end"])
                    if prev is None or it["filed"] > prev["filed"]:
                        by_end[it["end"]] = it
                if by_end:
                    table[label] = {end: v["val"] for end, v in by_end.items()}
                    break

        if not table:
            return None
        df = pd.DataFrame(table).T
        df = df[sorted(df.columns, reverse=True)[:years]]
        df.index.name = "Metric"
        return df


def style_sheet(ws, money_cols=None, link_col=None):
    header_fill = PatternFill("solid", start_color="1F3864")
    for cell in ws[1]:
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT)
            if money_cols and cell.column in money_cols and isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0;(#,##0);"-"'
            if link_col and cell.column == link_col and cell.value:
                cell.hyperlink = cell.value
                cell.value = "Open filing"
                cell.font = Font(name=FONT, color="0563C1", underline="single")

    for col in ws.columns:
        longest = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(longest + 2, 12), 45)
    ws.freeze_panes = "A2"


def main():
    p = argparse.ArgumentParser(description="Export SEC EDGAR data to Excel")
    p.add_argument("tickers", nargs="+", help="Stock tickers, e.g. AAPL MSFT")
    p.add_argument("--user-agent", default=os.environ.get("EDGAR_USER_AGENT"),
                   help='Required by the SEC, e.g. "Jane Doe jane@example.com"')
    p.add_argument("--forms", nargs="+", default=["10-K", "10-Q", "8-K"],
                   help="Filing types to include (default: 10-K 10-Q 8-K)")
    p.add_argument("--count", type=int, default=20,
                   help="Max filings per company (default 20)")
    p.add_argument("--years", type=int, default=5,
                   help="Years of annual financials (default 5)")
    p.add_argument("--no-financials", action="store_true",
                   help="Only export the filings list")
    p.add_argument("-o", "--output", default="edgar_data.xlsx", help="Output .xlsx file")
    args = p.parse_args()

    if not args.user_agent or "@" not in args.user_agent:
        sys.exit('Error: the SEC requires a contact email. Use --user-agent '
                 '"Your Name your@email.com" (or set EDGAR_USER_AGENT).')

    edgar = Edgar(args.user_agent)
    all_filings, fin_frames = [], {}

    for ticker in args.tickers:
        cik, name = edgar.lookup(ticker)
        if not cik:
            print(f"[skip] Ticker not found: {ticker}", file=sys.stderr)
            continue
        print(f"[ok] {ticker.upper()} -> {name} (CIK {cik})", file=sys.stderr)

        for row in edgar.filings(cik, set(args.forms), args.count):
            all_filings.append({"Ticker": ticker.upper(), "Company": name, **row})

        if not args.no_financials:
            df = edgar.annual_financials(cik, args.years)
            if df is None:
                print(f"[warn] No annual XBRL financials found for {ticker}", file=sys.stderr)
            else:
                fin_frames[ticker.upper()] = df

    if not all_filings and not fin_frames:
        sys.exit("Nothing to save.")

    with pd.ExcelWriter(args.output, engine="openpyxl") as writer:
        if all_filings:
            pd.DataFrame(all_filings).to_excel(writer, sheet_name="Filings", index=False)
        for ticker, df in fin_frames.items():
            df.to_excel(writer, sheet_name=f"{ticker} Financials")

    wb = load_workbook(args.output)
    if "Filings" in wb.sheetnames:
        ws = wb["Filings"]
        link_idx = [c.value for c in ws[1]].index("Link") + 1
        style_sheet(ws, link_col=link_idx)
    for ticker in fin_frames:
        ws = wb[f"{ticker} Financials"]
        style_sheet(ws, money_cols=set(range(2, ws.max_column + 1)))
        for row in ws.iter_rows(min_row=2):
            if row[0].value == "Diluted EPS":
                for cell in row[1:]:
                    cell.number_format = "0.00"
        note_row = ws.max_row + 2
        ws.cell(note_row, 1, "Source: SEC EDGAR XBRL company facts API (10-K annual data). "
                             "Values in USD; EPS in USD per share. "
                             "Blank = not reported under the tags this script checks.").font = \
            Font(name=FONT, italic=True, size=9)
    wb.save(args.output)
    print(f"Saved to {args.output}", file=sys.stderr)


if __name__ == "__main__":
    main()